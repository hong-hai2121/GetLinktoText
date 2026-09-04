"""
billing_core.py — Lõi xử lý dùng chung cho fb_ads_noApi (không có giao diện).

  • Đọc file xuất từ extension FB Billing Scraper (.xlsx / .xls / .csv)
  • Lọc "Đã thanh toán", lọc tháng đa số, so cột, upsert theo "ID giao dịch"
  • Kết nối Google Sheet bằng service_account.json nằm CÙNG thư mục
  • Tên tab đích = phần tên file trước "_fb_billing…" (CHI_TIET → "<tên> Chi tiết")
  • Tên sheet mới mặc định: "Báo cáo ADS tháng x" (x = tháng hiện tại - 1)

Không import config.py của fb_ads_accounts (file đó đòi FB token); chỉ đọc
SPREADSHEET_ID trong đó bằng regex làm gợi ý cho lựa chọn "Sheet có sẵn".
"""

import csv
import json
import re
import sys
import time
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import date
from pathlib import Path

try:
    import openpyxl
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False

import gspread
from gspread.exceptions import APIError, SpreadsheetNotFound
from google.oauth2.service_account import Credentials

# ── Cấu hình ───────────────────────────────────────────────────────────────────

BASE_DIR      = Path(__file__).parent
SA_FILE       = BASE_DIR / "service_account.json"
SETTINGS_FILE = Path.home() / ".fb_ads_noApi.json"   # nhớ email chia sẻ, sheet gần nhất
SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

ID_COL           = "ID giao dịch"
STATUS_COL       = "Trạng thái"
STATUS_ALLOW     = "Đã thanh toán"
DATE_COL         = "Ngày thanh toán"   # cột ngày để lọc theo tháng
KNOWN_MONEY_COLS = ["Số tiền", "Thuế", "Tổng tiền giao dịch", "Số tiền phân bổ"]
TEXT_ID_COLS     = [
    "ID giao dịch",
    "Tài khoản quảng cáo",
    "ID nhóm quảng cáo",
    "Số tham chiếu",
    "Số tham chiếu thẻ",
]

SHEETS_NEW_URL = "https://sheets.new"


# ── Google Sheet helpers ───────────────────────────────────────────────────────

def get_client():
    creds = Credentials.from_service_account_file(str(SA_FILE), scopes=SCOPES)
    return gspread.authorize(creds)


def sa_email() -> str:
    """Email của service account — người dùng phải chia sẻ sheet cho email này."""
    try:
        return json.loads(SA_FILE.read_text(encoding="utf-8")).get("client_email", "")
    except Exception:
        return ""


def default_sheet_id_from_config() -> str:
    """Đọc SPREADSHEET_ID trong config.py (thư mục này, rồi fb_ads_accounts) bằng regex — KHÔNG import."""
    for cfg in (BASE_DIR / "config.py", BASE_DIR.parent / "fb_ads_accounts" / "config.py"):
        try:
            txt = cfg.read_text(encoding="utf-8")
        except OSError:
            continue
        m = re.search(r"SPREADSHEET_ID\s*=\s*[\"']([A-Za-z0-9_-]+)[\"']", txt)
        if m:
            return m.group(1)
    return ""


def extract_sheet_id(text: str) -> str:
    """Nhận link đầy đủ hoặc ID trần → trả ID; chuỗi không hợp lệ → ''."""
    s = (text or "").strip()
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", s)
    if m:
        return m.group(1)
    return s if re.fullmatch(r"[A-Za-z0-9_-]{20,}", s) else ""


def is_quota_error(exc: Exception) -> bool:
    s = str(exc).lower()
    return "quota" in s or "storage" in s


def load_settings() -> dict:
    try:
        return json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_settings(data: dict):
    try:
        SETTINGS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                                 encoding="utf-8")
    except OSError:
        pass


def default_sheet_title(today: date = None) -> str:
    """Tên sheet mới mặc định: 'Báo cáo ADS tháng x', x = tháng hiện tại - 1 (tháng 1 → tháng 12)."""
    today = today or date.today()
    prev_month = 12 if today.month == 1 else today.month - 1
    return f"Báo cáo ADS tháng {prev_month}"


def tab_name_from_stem(stem: str) -> str:
    """Tên tab = phần TRƯỚC hậu tố '_fb_billing…', giữ nguyên '(ID)'.

    File CHI TIẾT có bộ cột khác hẳn file danh sách → tab riêng '<tên> Chi tiết'.
    Google Sheets cấm  [ ] * ? / \\ :  trong tên tab và giới hạn 100 ký tự.
    """
    m = re.match(r"^(.+?)_fb_billing", stem, re.IGNORECASE)
    name = (m.group(1) if m else stem).strip()
    if re.search(r"_fb_billing_chi[_ ]?tiet", stem, re.IGNORECASE):
        name += " Chi tiết"
    name = re.sub(r"[\[\]\*\?/\\:]", "_", name).strip()
    return name[:100] or "Sheet1"


# ── Đọc file ───────────────────────────────────────────────────────────────────

def parse_vnd(text: str):
    if not text or not isinstance(text, str):
        return None
    s = text.replace("₫", "").replace("'", "").strip()
    if re.search(r"\d\.\d{3}", s):
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s.strip()) if s.strip() else None
    except ValueError:
        return None


def _money_cell(raw):
    """Làm sạch 1 ô tiền → SỐ thật (int/float) để Sheet cộng/SUM được; không parse được → chuỗi."""
    s = "" if raw is None else str(raw)
    val = parse_vnd(s)
    if val is None:
        return s.lstrip("'").strip()
    return int(val) if float(val).is_integer() else val


def _clean_status(raw) -> str:
    """Extension đôi khi xuất 'Đã thanh toánĐã thanh toán' (chuỗi lặp đôi) → rút về 1 lần."""
    s = "" if raw is None else str(raw).strip()
    n = len(s)
    if n and n % 2 == 0 and s[: n // 2] == s[n // 2:]:
        return s[: n // 2]
    return s


def read_csv(path: Path):
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            with open(path, encoding=enc, newline="") as f:
                sample = f.read(8192)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',\t;|')
                delim = dialect.delimiter
            except csv.Error:
                delim = ','
            with open(path, encoding=enc, newline="") as f:
                reader = csv.DictReader(f, delimiter=delim)
                headers = list(reader.fieldnames or [])
                rows    = list(reader)
            return headers, rows
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"Không đọc được file: {path}")


def read_xls(path: Path):
    """Đọc file SpreadsheetML (.xls) do FB Billing Scraper bản cũ xuất ra."""
    NS = 'urn:schemas-microsoft-com:office:spreadsheet'
    root = ET.parse(path).getroot()
    worksheet = root.find(f'{{{NS}}}Worksheet')
    if worksheet is None:
        raise RuntimeError("Không tìm thấy Worksheet trong file XLS")
    table = worksheet.find(f'{{{NS}}}Table')
    if table is None:
        raise RuntimeError("Không tìm thấy Table trong file XLS")

    all_rows = []
    for row_el in table.findall(f'{{{NS}}}Row'):
        row = []
        for cell_el in row_el.findall(f'{{{NS}}}Cell'):
            data_el = cell_el.find(f'{{{NS}}}Data')
            row.append(data_el.text or '' if data_el is not None else '')
        all_rows.append(row)
    if not all_rows:
        return [], []
    headers = all_rows[0]
    rows = [dict(zip(headers, r + [''] * max(0, len(headers) - len(r))))
            for r in all_rows[1:]]
    return headers, rows


def _col_idx(ref: str) -> int:
    col = ''.join(c for c in ref if c.isalpha())
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n - 1


def _read_xlsx_manual(path: Path):
    """Parse ZIP local headers trực tiếp — bỏ qua central directory bị hỏng."""
    import struct as _struct, zlib as _zlib

    with open(path, 'rb') as f:
        data = f.read()

    files: dict = {}
    pos = 0
    while pos < len(data) - 4:
        if data[pos:pos+4] == b'PK\x03\x04':
            _, _, compress, _, _, _, comp_sz, _, fname_len, extra_len = \
                _struct.unpack_from('<HHHHH IIIHH', data, pos + 4)
            fname = data[pos+30:pos+30+fname_len].decode('utf-8', errors='replace')
            ds    = pos + 30 + fname_len + extra_len
            raw   = data[ds:ds+comp_sz]
            files[fname] = _zlib.decompress(raw, -15) if compress == 8 else raw
            pos = ds + comp_sz
        elif data[pos:pos+4] == b'PK\x01\x02':
            break
        else:
            pos += 1

    sheet_xml = files.get('xl/worksheets/sheet1.xml')
    if not sheet_xml:
        raise RuntimeError("Không tìm thấy xl/worksheets/sheet1.xml trong file XLSX")

    NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    shared: list = []
    ss_xml = files.get('xl/sharedStrings.xml')
    if ss_xml:
        for si in ET.fromstring(ss_xml).findall(f'{{{NS}}}si'):
            t_el = si.find(f'{{{NS}}}t')
            shared.append(t_el.text or '' if t_el is not None else '')

    def cell_val(c: ET.Element) -> str:
        t = c.get('t', '')
        if t == 'inlineStr':
            is_t = c.find(f'{{{NS}}}is/{{{NS}}}t')
            return (is_t.text or '') if is_t is not None else ''
        v_el = c.find(f'{{{NS}}}v')
        if v_el is None:
            return ''
        if t == 's' and shared:
            idx = int(v_el.text or 0)
            return shared[idx] if idx < len(shared) else ''
        return v_el.text or ''

    sheet_data = ET.fromstring(sheet_xml).find(f'{{{NS}}}sheetData')
    if sheet_data is None:
        return [], []
    all_rows: list = []
    for row_el in sheet_data.findall(f'{{{NS}}}row'):
        cells = row_el.findall(f'{{{NS}}}c')
        if not cells:
            continue
        max_col = max(_col_idx(c.get('r', 'A')) for c in cells)
        row = [''] * (max_col + 1)
        for c in cells:
            row[_col_idx(c.get('r', 'A'))] = cell_val(c)
        all_rows.append(row)
    if not all_rows:
        return [], []
    headers = all_rows[0]
    rows = [dict(zip(headers, r + [''] * max(0, len(headers) - len(r))))
            for r in all_rows[1:]]
    return headers, rows


def read_xlsx(path: Path):
    """Đọc .xlsx: openpyxl trước, fallback manual ZIP parser (extension xuất ZIP không chuẩn)."""
    if _OPENPYXL_OK:
        try:
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")   # "Workbook contains no default style"
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                wb.close()
                return [], []
            headers = [str(h) if h is not None else '' for h in header_row]
            rows = [
                {headers[i]: (str(v) if v is not None else '') for i, v in enumerate(raw)}
                for raw in rows_iter
            ]
            wb.close()
            return headers, rows
        except zipfile.BadZipFile:
            pass
    return _read_xlsx_manual(path)


def read_file(path: Path):
    ext = path.suffix.lower()
    if ext == '.xls':
        headers, rows = read_xls(path)
    elif ext == '.xlsx':
        headers, rows = read_xlsx(path)
    else:
        headers, rows = read_csv(path)
    if STATUS_COL in headers:
        for r in rows:
            r[STATUS_COL] = _clean_status(r.get(STATUS_COL, ""))
    return headers, rows


# ── Lọc / so cột / upsert (giữ nguyên logic bản cũ) ───────────────────────────

def _normalize(name: str) -> str:
    return re.sub(r"[^\w]", "", name, flags=re.UNICODE).lower()


def _find_col(headers: list, name: str):
    target = _normalize(name)
    return next((h for h in headers if _normalize(h) == target), None)


def filter_by_status(csv_headers: list, csv_rows: list) -> tuple:
    if STATUS_COL not in csv_headers:
        return csv_rows, 0
    kept    = [r for r in csv_rows if STATUS_ALLOW in str(r.get(STATUS_COL, ""))]
    return kept, len(csv_rows) - len(kept)


_EN_MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
              "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}


def _extract_month_key(text: str):
    """'26 Tháng 8, 2026' · '2026-06-01' · '01/06/2026' · 'Jun 1, 2026' → (năm, tháng)."""
    if not text:
        return None
    s = str(text).strip()
    m = re.search(r"\b(\d{4})[-/](\d{1,2})[-/]\d{1,2}\b", s)
    if m:
        return (int(m.group(1)), int(m.group(2)))
    m = re.search(r"\b(\d{1,2})[/](\d{1,2})[/](\d{4})\b", s)
    if m:
        return (int(m.group(3)), int(m.group(2)))
    mv = re.search(r"th(?:án)?g\.?\s*(\d{1,2})", s, re.IGNORECASE)
    if mv:
        my = re.search(r"\b(20\d{2})\b", s)
        return (int(my.group(1)) if my else None, int(mv.group(1)))
    me = re.search(r"\b([A-Za-z]{3,})\b", s)
    if me and me.group(1)[:3].lower() in _EN_MONTHS:
        my = re.search(r"\b(20\d{2})\b", s)
        return (int(my.group(1)) if my else None, _EN_MONTHS[me.group(1)[:3].lower()])
    return None


def _month_label(key) -> str:
    y, m = key
    return f"tháng {m}/{y}" if y else f"tháng {m}"


def filter_by_month(csv_headers: list, csv_rows: list, target=None) -> tuple:
    """Giữ dòng thuộc tháng đa số; dòng không đọc được ngày thì GIỮ. Trả (kept, dropped, target, unparsed)."""
    col = _find_col(csv_headers, DATE_COL)
    if not col:
        return csv_rows, [], None, 0
    keys   = [_extract_month_key(r.get(col, "")) for r in csv_rows]
    parsed = [k for k in keys if k]
    if not parsed:
        return csv_rows, [], None, len(csv_rows)
    if target is None:
        target = Counter(parsed).most_common(1)[0][0]
    kept, dropped, unparsed = [], [], 0
    for row, key in zip(csv_rows, keys):
        if key is None:
            kept.append(row)
            unparsed += 1
        elif key == target:
            kept.append(row)
        else:
            dropped.append((row, key))
    return kept, dropped, target, unparsed


def analyze_structure(csv_headers: list, sheet_headers: list) -> dict:
    def keep(h: str) -> bool:
        return bool(h.strip()) and not h.endswith(" (numeric)")
    clean_sheet = [h for h in sheet_headers if keep(h)]
    clean_csv   = [h for h in csv_headers   if keep(h)]
    csv_norm    = {_normalize(h): h for h in clean_csv}
    sheet_norm  = {_normalize(h): h for h in clean_sheet}
    return {
        "matched":     [h for h in clean_csv   if _normalize(h) in sheet_norm],
        "new_cols":    [h for h in clean_csv   if _normalize(h) not in sheet_norm],
        "missing":     [h for h in clean_sheet if _normalize(h) not in csv_norm],
        "has_id":      ID_COL in set(csv_headers),
        "sheet_empty": len(clean_sheet) == 0,
    }


def _format_text_columns(worksheet, headers: list):
    requests = []
    for col_name in TEXT_ID_COLS:
        if col_name not in headers:
            continue
        ci = headers.index(col_name)
        requests.append({"repeatCell": {
            "range": {"sheetId": worksheet.id,
                      "startColumnIndex": ci, "endColumnIndex": ci + 1},
            "cell": {"userEnteredFormat": {"numberFormat": {"type": "TEXT"}}},
            "fields": "userEnteredFormat.numberFormat",
        }})
    if requests:
        worksheet.spreadsheet.batch_update({"requests": requests})


def upsert(worksheet, csv_headers: list, csv_rows: list, replace: bool = False) -> dict:
    csv_norm_map = {_normalize(h): h for h in csv_headers}

    if replace:
        sheet_hdr, sheet_body = list(csv_headers), []
    else:
        existing = worksheet.get_all_values()
        sheet_hdr  = existing[0] if existing else []
        sheet_body = [list(r) for r in existing[1:]] if existing else []
        if not sheet_hdr:
            sheet_hdr, sheet_body = list(csv_headers), []
        else:
            sheet_norm = {_normalize(h) for h in sheet_hdr}
            for h in csv_headers:
                if _normalize(h) not in sheet_norm:
                    sheet_hdr.append(h)
                    sheet_norm.add(_normalize(h))
            ncols = len(sheet_hdr)
            for r in sheet_body:
                while len(r) < ncols:
                    r.append("")

    id_idx = sheet_hdr.index(ID_COL) if ID_COL in sheet_hdr else None

    def _get_csv_val(row: dict, col_name: str) -> str:
        if col_name in row:
            return str(row[col_name])
        norm = csv_norm_map.get(_normalize(col_name))
        return str(row[norm]) if norm and norm in row else ""

    # File CHI TIẾT: nhiều dòng cùng 1 ID giao dịch → dedup CẤP GIAO DỊCH
    incoming_ids = ([str(r.get(ID_COL, "")).strip() for r in csv_rows]
                    if ID_COL in csv_headers else [])
    incoming_ids = [i for i in incoming_ids if i]
    is_detail = (len(incoming_ids) != len(set(incoming_ids))
                 or _find_col(csv_headers, "Tên chiến dịch") is not None)

    if id_idx is not None and is_detail:
        id_set = set(incoming_ids)
        kept, updated = [], 0
        for row in sheet_body:
            rid = row[id_idx].strip() if id_idx < len(row) else ""
            if rid and rid in id_set:
                updated += 1
                continue
            kept.append(row)
        sheet_body = kept
        added = 0
        for csv_row in csv_rows:
            sheet_body.append([_get_csv_val(csv_row, h) for h in sheet_hdr])
            added += 1
    else:
        existing_ids: dict = {}
        if id_idx is not None:
            for i, row in enumerate(sheet_body):
                rid = row[id_idx].strip() if id_idx < len(row) else ""
                if rid:
                    existing_ids[rid] = i
        updated = added = 0
        for csv_row in csv_rows:
            row_vals = [_get_csv_val(csv_row, h) for h in sheet_hdr]
            csv_id = str(csv_row.get(ID_COL, "")).strip() if ID_COL in csv_headers else ""
            if csv_id and csv_id in existing_ids:
                sheet_body[existing_ids[csv_id]] = row_vals
                updated += 1
            else:
                if csv_id:
                    existing_ids[csv_id] = len(sheet_body)
                sheet_body.append(row_vals)
                added += 1

    money_norm = {_normalize(c) for c in KNOWN_MONEY_COLS}
    money_idx  = [i for i, h in enumerate(sheet_hdr) if _normalize(h) in money_norm]
    if money_idx:
        for row in sheet_body:
            for i in money_idx:
                if i < len(row):
                    row[i] = _money_cell(row[i])

    worksheet.clear()
    all_data = [sheet_hdr] + sheet_body
    BATCH = 1000
    for start in range(0, len(all_data), BATCH):
        worksheet.update(values=all_data[start:start + BATCH],
                         range_name=f"A{start + 1}", value_input_option="RAW")
    _format_text_columns(worksheet, sheet_hdr)
    return {"tab": worksheet.title, "total": len(csv_rows),
            "updated": updated, "added": added}


def upsert_with_retry(ws, headers, rows, replace, max_retry=4):
    """Thử lại khi Google trả 429 / 5xx (hay gặp khi import nhiều file liền nhau)."""
    delay = 5
    for attempt in range(max_retry + 1):
        try:
            return upsert(ws, headers, rows, replace=replace)
        except APIError as e:
            code = getattr(getattr(e, "response", None), "status_code", None)
            if code in (429, 500, 502, 503) and attempt < max_retry:
                time.sleep(delay)
                delay = min(delay * 2, 60)
                continue
            raise


# ── Báo cáo xuất tay từ Ads Manager (Nhóm quảng cáo / Chiến dịch / Quảng cáo) ──

SPEND_COL_HINT = "chi tiêu"          # cột "Số tiền đã chi tiêu (VND)"
REPORT_TAB_PREFIX = "Camp "          # tab báo cáo = "Camp <Tên TK>", phân biệt với tab billing "<Tên TK> (<ID>)"
_REPORT_LEVELS = [                    # (mẫu trong tên file, nhãn) — thử "Nhóm quảng cáo" TRƯỚC "Quảng cáo"
    (r"Nh.m-qu.ng-c.o|Ad-?sets?", "Nhóm quảng cáo"),      # "." thay chữ có dấu: khớp cả tên file bị mất dấu
    (r"Chi.n-d.ch|Campaigns?", "Chiến dịch"),             #   ("Nh_m-qu_ng-c_o", "Chi_n-d_ch")
    (r"Qu.ng-c.o|Ads", "Quảng cáo"),
]


def parse_report_stem(stem: str) -> dict:
    """Tên file Ads Manager: '<Tên TK, khoảng trắng → "-">-<Cấp>-<từ ngày>-<đến ngày>'.

    'Hiệp-01---BSH-Hưng-Việt---GMT+7-(THUẾ)-Nhóm-quảng-cáo-1-Tháng-8-2026-31-Tháng-8-2026'
      → account 'Hiệp 01 - BSH Hưng Việt - GMT+7 (THUẾ)', level 'Nhóm quảng cáo',
        period '1 Tháng 8 2026 31 Tháng 8 2026'
    """
    level, period, raw = "", "", stem
    for pat, label in _REPORT_LEVELS:
        m = re.search(rf"-(?:{pat})-(.*)$", stem, re.IGNORECASE)
        if m:
            level, period, raw = label, m.group(1), stem[:m.start()]
            break
    account = raw.replace("---", "\x00").replace("-", " ").replace("\x00", " - ")
    account = re.sub(r"\s+", " ", account).strip()
    return {"account": account, "account_raw": raw, "level": level,
            "period": period.replace("-", " ").strip()}


def report_tab_name(account: str) -> str:
    """Tên tab báo cáo Ads Manager cho tài khoản: 'Camp Hiệp 01 - BSH Hưng Việt - GMT+7 (THUẾ)'."""
    account = (account or "").strip()
    return account if account.casefold().startswith(REPORT_TAB_PREFIX.casefold()) else REPORT_TAB_PREFIX + account


def _norm_key(s: str) -> str:
    """Khóa so sánh mềm: bỏ khoảng trắng, gạch, dấu câu; giữ chữ, số, ngoặc, dấu +."""
    return re.sub(r"[^\w()+]", "", (s or "").casefold())


def _ascii_key(s: str) -> str:
    """Khóa so sánh thô: chỉ giữ chữ cái ASCII và số. Dùng khi tên file bị mất dấu tiếng Việt —
    'H_ ng H_i 01 - BSH H_ng Vi_t - GMT_7 _THU_' và 'Hồng Hải 01 - BSH Hưng Việt - GMT+7 (THUẾ)'
    cùng ra 'hnghi01bshhngvitgmt7thu'."""
    return re.sub(r"[^a-z0-9]", "", (s or "").casefold())


def match_account_tab(account: str, titles) -> str:
    """Tìm tab đã có trong sheet ứng với tài khoản này → trả tên đúng chính tả của tab.

    Ưu tiên tab billing '<tên> (<ID>)' (trả phần tên, bỏ ID) — đúng quy ước Tổng Hợp T6:
    tab chiến dịch = tên tài khoản không có ID. Sau đó mới tới tab trùng tên không ID.
    Không khớp chính xác → so phần ASCII còn lại (tên file mất dấu), chỉ nhận khi đúng 1 tab khớp.
    """
    key = _norm_key(account)
    if not key:
        return ""
    titles = list(titles)
    strip_id = lambda t: re.sub(r"\s*\(\d{6,}\)\s*$", "", t)
    for t in titles:
        base = strip_id(t)
        if base != t and _norm_key(base) == key:
            return base.strip()
    for t in titles:
        if _norm_key(t) == key:
            return t
    akey = _ascii_key(account)
    if len(akey) >= 6:
        hits = {strip_id(t).strip() for t in titles if strip_id(t) != t and _ascii_key(strip_id(t)) == akey}
        if len(hits) == 1:
            return hits.pop()
        hits = {t for t in titles if _ascii_key(t) == akey}
        if len(hits) == 1:
            return hits.pop()
    return ""


def read_ads_report(path: Path):
    """Đọc file báo cáo Ads Manager (.xlsx/.csv) → (headers, rows) — rows là list, số giữ là số."""
    if path.suffix.lower() == ".csv":
        headers, dict_rows = read_csv(path)
        rows = [[r.get(h, "") for h in headers] for r in dict_rows]
    else:
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            first = next(it)
        except StopIteration:
            wb.close()
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in first]
        rows = [list(r) for r in it if any(v not in (None, "") for v in r)]
        wb.close()
    while headers and not headers[-1]:          # bỏ cột trống cuối
        headers.pop()
    rows = [r[:len(headers)] + [None] * max(0, len(headers) - len(r)) for r in rows]
    return headers, rows


def find_spend_col(headers) -> int | None:
    return next((i for i, h in enumerate(headers) if SPEND_COL_HINT in str(h).casefold()), None)


def _spent(v) -> bool:
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    s = str(v).strip().replace("₫", "").replace("đ", "").replace(".", "").replace(",", "").strip()
    try:
        return float(s) != 0
    except ValueError:
        return False


def filter_spend_rows(headers, rows):
    """Chỉ giữ dòng có 'Số tiền đã chi tiêu' khác 0. Trả (kept, skipped, spend_idx)."""
    idx = find_spend_col(headers)
    if idx is None:
        return rows, 0, None
    kept = [r for r in rows if idx < len(r) and _spent(r[idx])]
    return kept, len(rows) - len(kept), idx


def sheet_cells(rows):
    """Chuẩn hóa ô để ghi lên Sheet: None → '', datetime → chuỗi ISO, số giữ nguyên."""
    out = []
    for r in rows:
        line = []
        for v in r:
            if v is None:
                line.append("")
            elif hasattr(v, "isoformat"):
                line.append(v.isoformat(sep=" ") if hasattr(v, "hour") else v.isoformat())
            else:
                line.append(v)
        out.append(line)
    return out
