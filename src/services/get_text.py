import gc
import os
import re
import sys
import threading
from typing import Callable

from docx import Document
from docx.shared import Pt, RGBColor

from src import browser_client
from src import config
from src import download
from src import io
from src import transcribe
from src.notifications import play_sound


class StopRequestedError(Exception):
    """Raised when user requests to stop current processing."""


class GetTextService:
    def __init__(self, logger: Callable[[str], None] | None = None):
        self.logger = logger or (lambda message: print(message))
        self._link_counter = 0
        self.max_links_per_run = 0
        self.processed_links_file = config.PROCESSED_LINKS_FILE
        self.failed_links_file = config.FAILED_LINKS_FILE
        self.move_processed_links = True
        self.stop_event = threading.Event()
        self._fail_reason = None

    def log(self, message: str):
        self.logger(message)

    def configure_runtime(
        self,
        model_name: str,
        use_title_mode: bool,
        title_submode: str,
        max_chunks: int,
        max_links_per_run: int = 0,
        skip_long_videos: bool = False,
        skip_gpt_processing: bool = False,
        move_processed_links: bool = True,
        language: str = "auto",
        srt_timestamps: bool = False,
        song_mode: bool = False,
    ):
        config.TRANSCRIBE_LANGUAGE = language or "auto"
        config.SRT_TIMESTAMPS = bool(srt_timestamps)
        config.SONG_MODE = bool(song_mode)
        config.USE_TITLE_MODE = use_title_mode
        config.TITLE_MODE_SUBMODE = title_submode
        config.MAX_CHUNKS_RUNTIME = max(1, int(max_chunks))
        config.SKIP_LONG_VIDEOS = bool(skip_long_videos)
        config.SKIP_GPT_PROCESSING = bool(skip_gpt_processing)
        config.MODEL_NAME = model_name
        self.max_links_per_run = max(0, int(max_links_per_run))
        self.processed_links_file = config.PROCESSED_LINKS_FILE
        self.failed_links_file = config.FAILED_LINKS_FILE
        self.move_processed_links = bool(move_processed_links)

    def initialize_whisper(self):
        # The distributable CPU edition does not require PyTorch or CUDA.
        _torch = None
        if not getattr(sys, "frozen", False):
            try:
                import torch as _torch
            except ImportError:
                pass
        from faster_whisper import WhisperModel as _WhisperModel

        config.torch = _torch
        config.WhisperModel = _WhisperModel
        config.device = "cuda" if _torch is not None and _torch.cuda.is_available() else "cpu"
        config.compute_type = "float16" if config.device == "cuda" else "int8"
        self.log(f"✅ Whisper sẽ chạy trên {'GPU' if config.device=='cuda' else 'CPU'} ({config.compute_type})")

    def load_model(self, selected_model: str):
        # Model giữ trong VRAM tới khi tắt app (xem config.loaded_models). Đổi model
        # thì nạp thêm chứ không xóa model cũ — xóa là crash.
        config.MODEL_NAME = selected_model
        cached = config.loaded_models.get(selected_model)
        if cached is not None:
            config.model = cached
            self.log(f"♻️ Dùng lại model '{selected_model}' đã nạp sẵn.")
            return

        self.log(f"🔄 Đang tải model '{selected_model}'...")
        config.model = config.WhisperModel(
            config.resolve_model(selected_model),
            device=config.device,
            compute_type=config.compute_type,
            download_root=config.WHISPER_MODEL_PATH,
        )
        config.loaded_models[selected_model] = config.model
        self.log(f"✅ Đã tải model '{selected_model}' thành công")

    def restart_firefox(self):
        return browser_client.restart_firefox()

    def request_stop(self):
        self.stop_event.set()

    def clear_stop_request(self):
        self.stop_event.clear()

    def is_stop_requested(self):
        return self.stop_event.is_set()

    def _check_stop_requested(self):
        if self.is_stop_requested():
            raise StopRequestedError("Đã nhận yêu cầu dừng từ người dùng.")

    def _read_links_file(self):
        with open(config.LINKS_FILE, "r", encoding="utf-8") as file:
            return [line.strip() for line in file.readlines() if line.strip()]

    def _split_links_for_current_run(self, links: list[str]):
        if self.max_links_per_run <= 0:
            return links, []
        return links[: self.max_links_per_run], links[self.max_links_per_run :]

    def _append_processed_links(self, processed_links: list[str]):
        if not processed_links:
            return

        target_dir = os.path.dirname(os.path.abspath(self.processed_links_file))
        if target_dir:
            os.makedirs(target_dir, exist_ok=True)

        with open(self.processed_links_file, "a", encoding="utf-8") as file:
            if file.tell() > 0:
                file.write("\n")
            file.write("\n".join(processed_links))
            file.write("\n")

    def _rewrite_remaining_links(self, remaining_links: list[str]):
        with open(config.LINKS_FILE, "w", encoding="utf-8") as file:
            if remaining_links:
                file.write("\n".join(remaining_links))
                file.write("\n")

    def _finalize_link_batch(self, processed_links: list[str], remaining_links: list[str]):
        if self.max_links_per_run <= 0 or not processed_links:
            return
        if not self.move_processed_links:
            self.log("Batch đã xử lý xong. Đã tắt chế độ chuyển link sang processed_links.txt.")
            return

        try:
            self._append_processed_links(processed_links)
            self._rewrite_remaining_links(remaining_links)
            self.log(
                f"✅ Đã chuyển {len(processed_links)} link đầu vào {self.processed_links_file} và cập nhật lại {config.LINKS_FILE}."
            )
        except Exception as e:
            self.log(f"⚠️ Không thể cập nhật file link sau khi xử lý batch: {e}")

    # ── Link chưa lấy được (failed_links.txt) ──────────────────────────────
    def reset_failed_links(self):
        """Xóa danh sách link lỗi cũ để tab 'Link chưa lấy' chỉ hiện lỗi của lần chạy này."""
        try:
            cu = io.read_failed_lines(self.failed_links_file)
            if not cu:
                return 0
            io.write_failed_lines(self.failed_links_file, [])
            self.log(f"🧹 Đã xóa {len(cu)} link chưa lấy của lần chạy trước.")
            return len(cu)
        except Exception as e:
            self.log(f"⚠️ Không xóa được danh sách link chưa lấy cũ: {e}")
            return 0

    def _mark_link_failed(self, link_key: str, reason: str):
        try:
            io.upsert_failed_link(self.failed_links_file, link_key, reason)
            self.log(f"⚠️ Chưa lấy được ({reason}) — đã ghi vào tab 'Link chưa lấy': {link_key}")
        except Exception as e:
            self.log(f"⚠️ Không ghi được link lỗi vào {self.failed_links_file}: {e}")

    def _mark_link_succeeded(self, link_key: str):
        # Link từng lỗi nay lấy được rồi -> tự bỏ khỏi danh sách link chưa lấy
        try:
            io.remove_failed_links(self.failed_links_file, {link_key})
        except Exception:
            pass

    def clean_link(self, line: str, fallback_index: int | None = None):
        """Tách (số thứ tự, link, nhãn) từ một dòng danh sách.

        Nhãn = phần chữ đứng trước https://... (ví dụ 'Kb25:'), được ghi SAU số
        thứ tự tự đánh trong file Word. Dòng chỉ có link trần thì nhãn = None.
        """
        line = line.strip()
        match = re.match(r"^\s*(\d+)[\.\s]+\s*(https?://[^\s\"'<>]+)", line)
        if match:
            return int(match.group(1)), match.group(2).rstrip(".,;)]}>"), None

        url_match = re.search(r"(https?://[^\s\"'<>]+)", line)
        if url_match:
            extracted_url = url_match.group(1).rstrip(".,;)]}>")
            label = line[: url_match.start()].strip().rstrip(":-–—.").strip() or None
            if fallback_index is not None:
                return fallback_index, extracted_url, label

            self._link_counter += 1
            return self._link_counter, extracted_url, label

        if fallback_index is not None:
            return fallback_index, line, None

        self._link_counter += 1
        return self._link_counter, line, None

    def check_title_blacklist(self, title: str):
        if not title or title == "Không thể lấy tiêu đề":
            return False

        title_clean = title.split("#")[0].strip() if "#" in title else title
        title_lower = title_clean.lower()

        for keyword in config.BLACKLIST_KEYWORDS:
            keyword_lower = keyword.lower()
            if keyword_lower in title_lower:
                self.log(
                    f"⚠️ Tiêu đề '{title}' chứa từ khóa cấm '{keyword}' (phần kiểm tra: '{title_clean}'), bỏ qua video."
                )
                return True
        return False

    def save_refined_text(self, refined_text: str | None, video_number: int):
        refined_text_filename = f"refined_text_{video_number}.txt"
        refined_text_path = os.path.join(config.OUTPUT_VIDEO_DIR, refined_text_filename)

        with open(refined_text_path, "w", encoding="utf-8") as f:
            if refined_text is None:
                self.log("⚠ Không có refined_text, bỏ qua xử lý văn bản.")
                refined_text = ""
            formatted_refined_text = refined_text.replace(". ", ".\n")
            f.write(formatted_refined_text)

    def download_video(self, link: str, video_number: int, browser):
        self._check_stop_requested()
        is_local_file = os.path.exists(link) and os.path.isfile(link)
        if is_local_file:
            self.log(f"📁 Phát hiện file local: {link}")
            title = os.path.splitext(os.path.basename(link))[0]
            video_path = link
            audio_filename = f"audio_{video_number}.wav"
            audio_path = os.path.join(config.OUTPUT_VIDEO_DIR, audio_filename)
            return title, video_path, audio_path, is_local_file

        checklinktiktok = 1 if download.is_tiktok_link(link) else 0
        self.log(f"{link} -> TikTok: {checklinktiktok}")

        title = download.get_video_title_facebook(link, browser)
        if config.CHECK_TITLE_BLACKLIST and self.check_title_blacklist(title):
            self.log(f"⛔ Bỏ qua video {video_number} do tiêu đề không phù hợp.")
            self._fail_reason = "Tiêu đề chứa từ khóa cấm"
            return None

        if config.SKIP_LONG_VIDEOS:
            duration_seconds = download.get_video_duration_seconds(link)
            if duration_seconds is not None and duration_seconds > config.LONG_VIDEO_THRESHOLD_SECONDS:
                max_minutes = config.LONG_VIDEO_THRESHOLD_SECONDS // 60
                self.log(
                    f"⏭️ Bo qua video {video_number}: thoi luong {duration_seconds / 60:.1f} phut > {max_minutes} phut."
                )
                self._fail_reason = f"Video dài {duration_seconds / 60:.1f} phút (đang bật bỏ qua video dài hơn {max_minutes} phút)"
                return None

        fb_video_url = None
        if checklinktiktok == 0:
            fb_video_url = download.get_facebook_video_url(link)
            if not fb_video_url:
                self._fail_reason = "Không lấy được URL video từ Facebook"
                return None

        video_filename = f"video_{video_number}.mp4"
        video_path = os.path.join(config.OUTPUT_VIDEO_DIR, video_filename)

        if checklinktiktok == 0:
            self.log("\n🎥 ĐANG TẢI VIDEO LINK FACEBOOK")
            if not download.download_video(fb_video_url, video_path):
                self._fail_reason = "Tải video Facebook thất bại"
                return None

        if checklinktiktok == 1:
            self.log("\n🎥 ĐANG TẢI VIDEO LINK TIKTOK")
            if not download.download_videoTIKTOK(link, video_path):
                self._fail_reason = "Tải video TikTok thất bại"
                return None

        audio_filename = f"audio_{video_number}.wav"
        audio_path = os.path.join(config.OUTPUT_VIDEO_DIR, audio_filename)
        return title, video_path, audio_path, is_local_file

    def transcribe_audio(self, video_path: str, audio_path: str, video_number: int):
        self._check_stop_requested()
        if os.path.exists(video_path):
            self.log(f"🎥 Video đã sẵn sàng: {video_path}")
        else:
            self.log(f"❌ Lỗi! Video chưa được sẵn sàng tại {video_path}")
            self._fail_reason = "Không tìm thấy file video đã tải"
            return None, None, None

        if not transcribe.extract_audio(video_path, audio_path):
            self._fail_reason = "Tách audio từ video thất bại (ffmpeg)"
            return None, None, None

        if os.path.exists(audio_path):
            self.log(f"🎵 Âm thanh đã lưu tại: {audio_path}")
        else:
            self.log(f"❌ Lỗi! Âm thanh chưa được lưu vào {audio_path}")
            self._fail_reason = "Không lưu được file audio"
            return None, None, None

        # Lấy thời lượng audio một lần để vừa kiểm tra video dài, vừa ghi vào Word
        duration_seconds = transcribe.get_audio_duration(audio_path)

        if config.SKIP_LONG_VIDEOS:
            if duration_seconds is None:
                self.log("⚠️ Khong lay duoc thoi luong audio, tiep tuc xu ly.")
            elif duration_seconds > config.LONG_VIDEO_THRESHOLD_SECONDS:
                max_minutes = config.LONG_VIDEO_THRESHOLD_SECONDS // 60
                self.log(
                    f"⏭️ Bo qua video {video_number}: thoi luong {duration_seconds / 60:.1f} phut > {max_minutes} phut."
                )
                self._fail_reason = f"Video dài {duration_seconds / 60:.1f} phút (đang bật bỏ qua video dài hơn {max_minutes} phút)"
                return None, None, duration_seconds

        transcript = transcribe.transcribe_whisper(
            audio_path,
            use_chunks=config.USE_TITLE_MODE,
            max_chunks=config.MAX_CHUNKS_RUNTIME,
            save_chunk_transcripts=config.SKIP_GPT_PROCESSING,
            chunk_text_prefix=os.path.join(config.OUTPUT_VIDEO_DIR, f"transcript_{video_number}"),
        )
        transcript_filename = f"transcript_{video_number}.txt"
        transcript_path = os.path.join(config.OUTPUT_VIDEO_DIR, transcript_filename)
        with open(transcript_path, "w", encoding="utf-8") as f:
            if transcript is None:
                self.log("⚠ Không có transcript, bỏ qua xử lý văn bản.")
                transcript = ""
            # SRT đã mỗi câu một dòng -> không ngắt dòng lại theo số từ
            if config.SRT_TIMESTAMPS:
                f.write(transcript)
            else:
                f.write(transcribe.format_transcript_by_spaces(transcript, spaces_per_line=25))

        self.log(f"💾 Đã lưu transcript vào: {transcript_path}")
        return transcript, transcript_path, duration_seconds

    @staticmethod
    def _format_duration_label(duration_seconds: float | None) -> str:
        """Trả về ' (video dài X phút Y giây)' từ thời lượng audio (giây); '' nếu không có."""
        if not duration_seconds or duration_seconds <= 0:
            return ""
        total = int(round(duration_seconds))
        if total <= 0:
            return ""
        minutes, seconds = divmod(total, 60)
        if minutes and seconds:
            text = f"{minutes} phút {seconds} giây"
        elif minutes:
            text = f"{minutes} phút"
        else:
            text = f"{seconds} giây"
        return f" (video dài {text})"

    def export_word(
        self,
        doc: Document,
        video_number: int,
        link: str,
        title: str,
        transcript: str | None,
        refined_text: str | None,
        title_suggestion: str | None = None,
        audio_duration: float | None = None,
        label: str | None = None,
    ):
        para = doc.add_paragraph()
        heading = f"{video_number} {label}. {link}" if label else f"{video_number}. {link}"
        run = para.add_run(heading)
        run.font.color.rgb = RGBColor(255, 0, 0)
        run.font.size = Pt(11)
        doc.add_paragraph(f"Tiêu đề: {title}")
        if title_suggestion:
            doc.add_paragraph(f"Tiêu đề gợi ý: {title_suggestion}")

        duration_label = self._format_duration_label(audio_duration)
        if transcript:
            if refined_text:
                doc.add_paragraph(f"Nội dung{duration_label}:\n{refined_text}")
            elif config.SKIP_GPT_PROCESSING or not config.USE_TITLE_MODE or (
                config.USE_TITLE_MODE and config.TITLE_MODE_SUBMODE == "content_only"
            ):
                doc.add_paragraph(f"Nội dung{duration_label}:\n{transcript}")
        else:
            doc.add_paragraph(f"Nội dung{duration_label}: (Video quá dài để lấy nội dung)")

        doc.add_paragraph("=" * 40)

        # title_only/full với GPT: caller thêm GPT rồi mới save — tránh double save
        # content_only và các chế độ khác: save ngay tại đây
        skip_save = (
            config.USE_TITLE_MODE
            and not config.SKIP_GPT_PROCESSING
            and config.TITLE_MODE_SUBMODE != "content_only"
        )
        if not skip_save:
            try:
                doc.save(config.OUTPUT_WORD_FILE)
                self.log("💾 Đã lưu kết quả thành công!")
            except Exception as e:
                self.log(f"❌ Lỗi khi lưu file Word: {e}")
        self.log("✅ Hoàn thành xử lý video, tiếp tục với video tiếp theo!")

    def process_video(self, index: int, link: str, doc: Document, driver, browser):
        self._check_stop_requested()
        self._fail_reason = None
        video_number, link, label = self.clean_link(link, index)
        link = link.strip().strip('"').strip("'")
        video_number = video_number if video_number is not None else index
        display_name = f"{video_number} {label}" if label else str(video_number)
        self.log(f"\n🎥 Đang xử lý video {display_name}: {link}")

        download_result = self.download_video(link, video_number, browser)
        if not download_result:
            return None
        title, video_path, audio_path, is_local_file = download_result

        transcript, transcript_path, audio_duration = self.transcribe_audio(video_path, audio_path, video_number)

        try:
            if os.path.exists(audio_path):
                os.remove(audio_path)
                self.log(f"🗑️ Đã xóa file âm thanh: {audio_path}")
        except Exception as e:
            self.log(f"⚠️ Không thể xóa audio: {e}")

        if not is_local_file and os.path.exists(video_path):
            try:
                os.remove(video_path)
                self.log(f"🗑️ Đã xóa file video: {video_path}")
            except Exception as e:
                self.log(f"⚠️ Không thể xóa video: {e}")

        if transcript is None:
            return transcript_path

        gc.collect()
        if config.device == "cuda":
            config.torch.cuda.empty_cache()

        if config.SKIP_GPT_PROCESSING:
            self.log(f"Video {video_number}: da bat che do chi luu transcript, bo qua GPT.")
            refined_text = None
            title_suggestion = None
        elif config.USE_TITLE_MODE:
            self.log("📁 Chế độ chia nhỏ: sẽ xử lý GPT sau khi nhận diện xong tất cả...")
            refined_text = None
            title_suggestion = None
        elif transcript:
            self.log(f"📝 Đang gửi nội dung video {video_number} đến ChatGPT...")
            refined_text = browser_client.send_to_chatgpt(driver, transcript)
            self.save_refined_text(refined_text, video_number)
            title_suggestion = None
        else:
            self.log(f"⚠️ Video {video_number} không có transcript, bỏ qua gửi đến ChatGPT.")
            refined_text = "Đây là video quá dài để lấy nội dung."
            title_suggestion = None

        self.export_word(
            doc, video_number, link, title, transcript, refined_text, title_suggestion, audio_duration, label=label
        )
        return transcript_path

    def process_videos(self):
        self._check_stop_requested()
        # Mỗi đợt chạy chỉ theo dõi link lỗi của chính nó (GUI đã xóa trước khi
        # bấm 'Bắt đầu'; gọi lại ở đây để các lối chạy khác cũng sạch danh sách cũ).
        self.reset_failed_links()
        doc = Document()
        doc.save(config.OUTPUT_WORD_FILE)

        driver = browser_client.global_driver

        need_gpt = (not config.SKIP_GPT_PROCESSING) and (
            (not config.USE_TITLE_MODE) or (config.USE_TITLE_MODE and config.TITLE_MODE_SUBMODE != "content_only")
        )

        if need_gpt:
            with browser_client.DRIVER_LOCK:
                in_progress = browser_client.FIREFOX_INIT_IN_PROGRESS

            if (driver is None or not browser_client.is_driver_alive(driver)) and in_progress:
                self.log("⏳ Firefox đang khởi động (thread nền). Đang chờ...")
                browser_client.FIREFOX_INIT_EVENT.wait(timeout=90)
                driver = browser_client.global_driver

            if driver is None or not browser_client.is_driver_alive(driver):
                self.log("🌐 Chưa có Firefox/driver. Tự mở Firefox để dùng GPT...")
                driver = browser_client.init_firefox()
                browser_client.global_driver = driver

        links = self._read_links_file()
        links_for_run, remaining_links = self._split_links_for_current_run(links)
        succeeded_lines: list[str] = []
        failed_lines: list[str] = []

        if self.max_links_per_run > 0:
            self.log(
                f"🎯 Dang chay batch {len(links_for_run)} link dau (gioi han={self.max_links_per_run}, tong={len(links)})."
            )

        for idx, raw_line in enumerate(links_for_run, start=1):
            self._check_stop_requested()
            video_number, link_key, _label = self.clean_link(raw_line)
            link_key = link_key.strip().strip('"').strip("'")
            self.log(f"🔄 Đang xử lý video {idx}/{len(links_for_run)} (ID={video_number})...")

            active_driver = driver if need_gpt else None
            try:
                transcript_file = self.process_video(idx, raw_line, doc, active_driver, active_driver)
            except StopRequestedError:
                raise
            except Exception as e:
                # Một link hỏng không làm đổ cả đợt chạy — ghi vào danh sách lỗi rồi đi tiếp
                self.log(f"❌ Lỗi khi xử lý {link_key}: {e}")
                self._fail_reason = f"Lỗi: {e}"
                transcript_file = None
            self._check_stop_requested()

            if (
                config.USE_TITLE_MODE
                and not config.SKIP_GPT_PROCESSING
                and transcript_file
                and os.path.exists(transcript_file)
            ):
                with open(transcript_file, "r", encoding="utf-8") as f:
                    current_transcript = f.read().strip()

                if current_transcript:
                    self._check_stop_requested()
                    if config.TITLE_MODE_SUBMODE == "content_only":
                        self.log(f"\n📝 Chế độ Nội dung: lưu nội dung video {idx} vào Word...")
                        try:
                            doc.save(config.OUTPUT_WORD_FILE)
                            self.log(f"💾 Đã lưu kết quả vào {config.OUTPUT_WORD_FILE}")
                        except Exception as e:
                            self.log(f"❌ Lỗi khi lưu file Word: {e}")
                        try:
                            os.remove(transcript_file)
                            self.log(f"🗑️ Đã xóa file transcript: {transcript_file}")
                        except Exception as e:
                            self.log(f"⚠️ Không thể xóa transcript: {e}")
                        succeeded_lines.append(raw_line)
                        self._mark_link_succeeded(link_key)
                        continue

                    self.log(f"\n📝 Đang xử lý GPT cho video {idx}...")
                    self.log("🔄 Sử dụng Firefox đã được mở từ khi khởi động...")

                    self.log(f"\n🎯 Yêu cầu 1: Gợi ý 5 tiêu đề cho video {idx}...")
                    title_prompt = (
                        "Dựa trên nội dung video sau, hãy đề xuất 5 tiêu đề hấp dẫn, phù hợp cho kênh Haio6VietSub. "
                        "Mỗi tiêu đề trên một dòng, không ghi số thứ tự:\n\n" + current_transcript
                    )
                    titles_result = browser_client.send_to_chatgpt(driver, title_prompt, mode="raw")
                    self._check_stop_requested()
                    if titles_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"🎯 5 Tiêu đề gợi ý cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(titles_result)

                    if config.TITLE_MODE_SUBMODE == "title_only":
                        doc.save(config.OUTPUT_WORD_FILE)
                        try:
                            os.remove(transcript_file)
                            self.log(f"🗑️ Đã xóa file transcript: {transcript_file}")
                        except Exception as e:
                            self.log(f"⚠️ Không thể xóa transcript: {e}")
                        self.log(f"✅ Video {idx}: Chế độ Tiêu đề - đã xong sau Yêu cầu 1.")
                        succeeded_lines.append(raw_line)
                        self._mark_link_succeeded(link_key)
                        continue

                    self.log(f"\n📝 Yêu cầu 2: Viết mô tả video {idx} và hashtag...")
                    desc_prompt = (
                        "Viết mô tả video (1-2 đoạn ngắn gọn) và thêm 5-10 hashtag phù hợp cho kênh Haio6VietSub, "
                        "dựa trên nội dung:\n\n" + current_transcript
                    )
                    desc_result = browser_client.send_to_chatgpt(driver, desc_prompt, mode="raw")
                    self._check_stop_requested()
                    if desc_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"📝 Mô tả + Hashtag cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(desc_result)

                    self.log(f"\n🏷️ Yêu cầu 3: Viết thẻ tag cho video {idx} (tối đa 500 kí tự)...")
                    tag_prompt = (
                        "Viết một danh sách các thẻ tag mô tả video này, tối đa 500 kí tự, phù hợp cho kênh Haio6VietSub, "
                        "ngăn cách bằng dấu phẩy. Dựa trên nội dung:\n\n" + current_transcript
                    )
                    tag_result = browser_client.send_to_chatgpt(driver, tag_prompt, mode="raw")
                    self._check_stop_requested()
                    if tag_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"🏷️ Tags cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(tag_result)

                    doc.save(config.OUTPUT_WORD_FILE)

                    try:
                        os.remove(transcript_file)
                        self.log(f"🗑️ Đã xóa file transcript: {transcript_file}")
                    except Exception as e:
                        self.log(f"⚠️ Không thể xóa transcript: {e}")

                    self.log(f"✅ Đã lưu kết quả video {idx} vào {config.OUTPUT_WORD_FILE}")

                    gc.collect()
                    if config.device == "cuda":
                        config.torch.cuda.empty_cache()

            if transcript_file is not None:
                succeeded_lines.append(raw_line)
                self._mark_link_succeeded(link_key)
            else:
                failed_lines.append(raw_line)
                self._mark_link_failed(link_key, self._fail_reason or "Không lấy được nội dung")

        # Link lỗi đã nằm trong failed_links.txt (lấy lại từ tab 'Link chưa lấy'),
        # nên không giữ lại trong danh sách chờ — tránh batch sau vấp lại đúng link hỏng.
        self._finalize_link_batch(succeeded_lines, remaining_links)

        # Không hủy model ở đây (xem load_model): chỉ dọn bộ nhớ tạm, model giữ lại
        # cho lần chạy sau và chỉ giải phóng khi tắt app.
        self.log("\n🧹 Dọn bộ nhớ tạm, giữ model cho lần chạy sau...")
        if config.device == "cuda":
            config.torch.cuda.empty_cache()
        gc.collect()
        self.log("✅ Xong. Model vẫn trong bộ nhớ, lần chạy sau không phải tải lại.")
        self.log("🎉 Hoàn thành tất cả video! Kết quả đã lưu vào final_result.docx")
        if failed_lines:
            self.log(
                f"⚠️ Có {len(failed_lines)} link chưa lấy được nội dung — mở tab '⛔ Link chưa lấy' để xem lý do và đưa về danh sách chạy lại."
            )
        play_sound("success")
