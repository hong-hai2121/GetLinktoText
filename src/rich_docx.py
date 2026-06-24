"""Đọc .docx giữ định dạng (màu/đậm/cỡ chữ) và copy sang clipboard dạng RTF.

GUI chỉ đọc plain text nên mất màu. Module này đọc theo từng *run* để giữ màu,
dựng chuỗi RTF và đẩy cả RTF lẫn plain text lên Windows clipboard => dán vào
Word/WordPad sẽ giữ nguyên màu & định dạng như file docx gốc.
"""

import os

from docx import Document


def read_docx_paragraphs(path):
    """Trả về list các đoạn; mỗi đoạn là list run dict(text, rgb, bold, size_pt)."""
    doc = Document(path)
    paragraphs = []
    for p in doc.paragraphs:
        runs = []
        for r in p.runs:
            text = r.text or ""
            if not text:
                continue

            rgb = None
            try:
                color = r.font.color
                if color is not None and color.rgb is not None:
                    hex_val = str(color.rgb)  # vd 'FF0000'
                    rgb = (
                        int(hex_val[0:2], 16),
                        int(hex_val[2:4], 16),
                        int(hex_val[4:6], 16),
                    )
            except Exception:
                rgb = None

            bold = bool(r.bold) if r.bold is not None else False

            size_pt = None
            try:
                if r.font.size is not None:
                    size_pt = float(r.font.size.pt)
            except Exception:
                size_pt = None

            runs.append({"text": text, "rgb": rgb, "bold": bold, "size_pt": size_pt})
        paragraphs.append(runs)
    return paragraphs


def read_docx_paragraphs_from_candidates(candidates):
    path = next((p for p in candidates if p and os.path.exists(p)), None)
    if not path:
        return None, []
    return path, read_docx_paragraphs(path)


def paragraphs_to_plain(paragraphs):
    lines = []
    for para in paragraphs:
        lines.append("".join(run.get("text", "") for run in para))
    return "\n".join(lines)


def paragraphs_to_docx(paragraphs, path):
    """Ghi các đoạn (giữ màu/đậm/cỡ chữ) ra file .docx mới và trả về đường dẫn."""
    from docx.shared import Pt, RGBColor

    doc = Document()
    for para in paragraphs:
        p = doc.add_paragraph()
        for run in para:
            r = p.add_run(run.get("text", ""))
            rgb = run.get("rgb")
            if rgb:
                r.font.color.rgb = RGBColor(*rgb)
            if run.get("bold"):
                r.bold = True
            size = run.get("size_pt")
            if size:
                r.font.size = Pt(size)
    doc.save(path)
    return path


def _rtf_escape(text):
    out = []
    for ch in text:
        cp = ord(ch)
        if ch == "\\":
            out.append("\\\\")
        elif ch == "{":
            out.append("\\{")
        elif ch == "}":
            out.append("\\}")
        elif ch == "\n":
            out.append("\\line ")
        elif ch == "\t":
            out.append("\\tab ")
        elif 32 <= cp < 127:
            out.append(ch)
        elif cp <= 0xFFFF:
            signed = cp - 65536 if cp > 32767 else cp
            out.append(f"\\u{signed}?")
        else:
            # ngoài BMP (emoji...) -> cặp surrogate, mỗi nửa 1 lệnh \u
            cp -= 0x10000
            for s in (0xD800 + (cp >> 10), 0xDC00 + (cp & 0x3FF)):
                signed = s - 65536 if s > 32767 else s
                out.append(f"\\u{signed}?")
    return "".join(out)


def paragraphs_to_rtf(paragraphs):
    """Dựng tài liệu RTF (ASCII thuần, unicode đã escape) từ các đoạn đã đọc."""
    colors = []
    color_index = {}

    def color_idx(rgb):
        if not rgb:
            return 0  # 0 = màu mặc định (đen/auto)
        if rgb not in color_index:
            colors.append(rgb)
            color_index[rgb] = len(colors)  # 1-based theo colortbl
        return color_index[rgb]

    body = []
    for para in paragraphs:
        for run in para:
            idx = color_idx(run.get("rgb"))
            bold = run.get("bold")
            size = run.get("size_pt")

            body.append(f"\\cf{idx} ")
            if size:
                body.append(f"\\fs{int(round(size * 2))} ")
            if bold:
                body.append("\\b ")
            body.append(_rtf_escape(run.get("text", "")))
            if bold:
                body.append("\\b0 ")
        body.append("\\par\n")

    colortbl = (
        "{\\colortbl;"
        + "".join(f"\\red{r}\\green{g}\\blue{b};" for (r, g, b) in colors)
        + "}"
    )
    header = (
        "{\\rtf1\\ansi\\ansicpg1252\\deff0"
        "{\\fonttbl{\\f0 Calibri;}}"
        + colortbl
        + "\\f0\\fs22\n"
    )
    return header + "".join(body) + "}"


def _html_escape(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def paragraphs_to_html(paragraphs):
    """Dựng HTML inline-style từ các đoạn đã đọc (cho Google Docs / trình duyệt)."""
    parts = []
    for para in paragraphs:
        if not para:
            parts.append("<p>&nbsp;</p>")
            continue
        spans = []
        for run in para:
            styles = []
            rgb = run.get("rgb")
            if rgb:
                styles.append("color:#%02X%02X%02X" % rgb)
            if run.get("bold"):
                styles.append("font-weight:bold")
            size = run.get("size_pt")
            if size:
                styles.append(f"font-size:{int(round(size))}pt")

            text = _html_escape(run.get("text", "")).replace("\n", "<br>")
            if styles:
                spans.append(f'<span style="{";".join(styles)}">{text}</span>')
            else:
                spans.append(text)
        parts.append("<p style=\"margin:0\">" + "".join(spans) + "</p>")
    return "".join(parts)


def _wrap_cf_html(html_fragment):
    """Bọc fragment HTML theo chuẩn CF_HTML (header có offset byte) cho clipboard Windows."""
    header = (
        "Version:0.9\r\n"
        "StartHTML:{:010d}\r\n"
        "EndHTML:{:010d}\r\n"
        "StartFragment:{:010d}\r\n"
        "EndFragment:{:010d}\r\n"
    )
    pre = "<html><body>\r\n<!--StartFragment-->"
    post = "<!--EndFragment-->\r\n</body></html>"

    # Độ dài header cố định (các offset đều dùng :010d nên format giả vẫn đúng độ dài)
    header_len = len(header.format(0, 0, 0, 0).encode("utf-8"))
    start_html = header_len
    start_fragment = start_html + len(pre.encode("utf-8"))
    end_fragment = start_fragment + len(html_fragment.encode("utf-8"))
    end_html = end_fragment + len(post.encode("utf-8"))

    return (
        header.format(start_html, end_html, start_fragment, end_fragment)
        + pre
        + html_fragment
        + post
    )


def copy_rich_to_clipboard(rtf_text, plain_text, html_text=None):
    """Đẩy RTF (Word), HTML (Google Docs/trình duyệt) và plain text lên clipboard. Cần pywin32."""
    import win32clipboard

    cf_rtf = win32clipboard.RegisterClipboardFormat("Rich Text Format")
    cf_html = win32clipboard.RegisterClipboardFormat("HTML Format")
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardData(
            win32clipboard.CF_UNICODETEXT, plain_text
        )
        win32clipboard.SetClipboardData(cf_rtf, rtf_text.encode("latin-1", "replace"))
        if html_text:
            win32clipboard.SetClipboardData(
                cf_html, _wrap_cf_html(html_text).encode("utf-8")
            )
    finally:
        win32clipboard.CloseClipboard()
