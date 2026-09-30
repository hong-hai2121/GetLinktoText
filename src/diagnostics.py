"""Offline checks for the installed executable, invoked with --self-test REPORT."""
import json
import os
from pathlib import Path
import subprocess
import tempfile
import traceback
import wave


def run_self_test(report_path):
    report = {"ok": False, "checks": []}
    try:
        from src import config
        from src.services.get_text import GetTextService
        from docx import Document
        import tkinter as tk
        import win32clipboard
        import yt_dlp
        from src.gui import GetTextApp

        config.ensure_directories()
        report["data_dir"] = config.BASE_DIR
        report["ffmpeg"] = config.FFMPEG_PATH
        report["model"] = config.resolve_model(config.MODEL_NAME)
        for program in (config.FFMPEG_PATH, config.FFPROBE_PATH, config.GECKODRIVER_PATH):
            subprocess.run([program, "-version" if "ff" in Path(program).name else "--version"],
                           check=True, capture_output=True, timeout=30)
        report["checks"].append("bundled ffmpeg, ffprobe and geckodriver")
        with tempfile.TemporaryDirectory(prefix="self-test-", dir=config.BASE_DIR) as work:
            doc_path = Path(work) / "result.docx"
            doc = Document()
            doc.add_paragraph("Tiếng Việt: kiểm tra bộ cài")
            doc.save(doc_path)
            assert Document(doc_path).paragraphs[0].text == "Tiếng Việt: kiểm tra bộ cài"
            report["checks"].append("Word Unicode write/read")
            root = tk.Tk()
            root.withdraw()
            try:
                GetTextApp(root)
                root.update()
            finally:
                root.destroy()
            report["checks"].append("Tk GUI and clipboard module")
            assert sum(1 for _ in yt_dlp.list_extractor_classes()) > 100
            report["checks"].append("yt-dlp extractors")
            service = GetTextService()
            service.initialize_whisper()
            assert config.device == "cpu"
            service.load_model(config.MODEL_NAME)
            audio_path = Path(work) / "silence.wav"
            with wave.open(str(audio_path), "wb") as audio:
                audio.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
                audio.writeframes(b"\x00\x00" * 16000)
            segments, _ = config.model.transcribe(str(audio_path), language="vi", vad_filter=True)
            list(segments)
            report["checks"].append("bundled Whisper model, audio decoder and ONNX VAD")
            speech = os.environ.get("GETTEXT_TEST_AUDIO")
            if speech:
                config.OUTPUT_VIDEO_DIR = work
                config.OUTPUT_WORD_FILE = str(doc_path)
                config.TRANSCRIBE_LANGUAGE = "en"
                config.SKIP_GPT_PROCESSING = True
                transcript, _, duration = service.transcribe_audio(
                    speech, str(Path(work) / "converted.wav"), 1,
                )
                assert transcript, "Speech produced an empty transcript"
                service.export_word(Document(), 1, speech, "Installer test", transcript, None,
                                    audio_duration=duration)
                word_text = "\n".join(p.text for p in Document(doc_path).paragraphs)
                assert transcript in word_text
                report["transcript"] = transcript
                report["checks"].append("FFmpeg -> CPU speech transcription -> Word export")
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
