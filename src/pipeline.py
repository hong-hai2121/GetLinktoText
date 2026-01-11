import gc
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor

from src import config
from src import download
from src import transcribe
from src import browser_client
from src.notifications import play_sound


def clean_link(line: str, fallback_index: int | None = None):
    """
    Nếu dòng có dạng '12. https://...' -> trả (12, url)
    Nếu không có số thứ tự:
      - Nếu có fallback_index: dùng fallback_index
      - Nếu không: tự tăng bộ đếm nội bộ (giống cách cũ nhưng không dùng mutable default arg)
    """
    match = re.match(r"^\s*(\d+)[\.\s]+\s*(https?://\S+)", line)
    if match:
        return int(match.group(1)), match.group(2)

    if fallback_index is not None:
        return fallback_index, line.strip()

    if not hasattr(clean_link, "_counter"):
        clean_link._counter = 0
    clean_link._counter += 1
    return clean_link._counter, line.strip()


def check_title_blacklist(title):
    """Kiểm tra xem tiêu đề có chứa từ khóa trong danh sách cấm hay không, bỏ qua phần từ # đến cuối"""
    if not title or title == "Không thể lấy tiêu đề":
        return False

    title_clean = title.split("#")[0].strip() if "#" in title else title
    title_lower = title_clean.lower()

    for keyword in config.BLACKLIST_KEYWORDS:
        keyword_lower = keyword.lower()
        if keyword_lower in title_lower:
            print(
                f"⚠️ Tiêu đề '{title}' chứa từ khóa cấm '{keyword}' (phần kiểm tra: '{title_clean}'), bỏ qua video."
            )
            return True
    return False


def save_refined_text(refined_text, video_number):
    refined_text_filename = f"refined_text_{video_number}.txt"
    refined_text_path = os.path.join(config.OUTPUT_VIDEO_DIR, refined_text_filename)

    with open(refined_text_path, "w", encoding="utf-8") as f:
        if refined_text is None:
            print("⚠ Không có refined_text, bỏ qua xử lý văn bản.")
            refined_text = ""
        formatted_refined_text = refined_text.replace(". ", ".\n")
        f.write(formatted_refined_text)


def process_video(index, link, doc, driver, browser):
    video_number, link = clean_link(link, index)
    link = link.strip().strip('"').strip("'")
    video_number = video_number if video_number is not None else index
    print(f"\n🎥 Đang xử lý video {video_number}: {link}")

    is_local_file = os.path.exists(link) and os.path.isfile(link)

    if is_local_file:
        print(f"📁 Phát hiện file local: {link}")
        title = os.path.splitext(os.path.basename(link))[0]
        video_path = link
        audio_filename = f"audio_{video_number}.wav"
        audio_path = os.path.join(config.OUTPUT_VIDEO_DIR, audio_filename)
    else:
        checklinktiktok = 1 if download.is_tiktok_link(link) else 0
        print(f"{link} -> TikTok: {checklinktiktok}")

        if checklinktiktok == 0:
            title = download.get_video_title_facebook(link, browser)
        else:
            title = download.get_video_title_facebook(link, browser)

        if config.CHECK_TITLE_BLACKLIST and check_title_blacklist(title):
            print(f"⛔ Bỏ qua video {video_number} do tiêu đề không phù hợp.")
            return

        if checklinktiktok == 0:
            fb_video_url = download.get_facebook_video_url(link)
            if not fb_video_url:
                return None

        video_filename = f"video_{video_number}.mp4"
        video_path = os.path.join(config.OUTPUT_VIDEO_DIR, video_filename)

        if checklinktiktok == 0:
            print("\n🎥 ĐANG TẢI VIDEO LINK FACEBOOK")
            if not download.download_video(fb_video_url, video_path):
                return None

        if checklinktiktok == 1:
            print("\n🎥 ĐANG TẢI VIDEO LINK TIKTOK")
            if not download.download_videoTIKTOK(link, video_path):
                return None
        audio_filename = f"audio_{video_number}.wav"
        audio_path = os.path.join(config.OUTPUT_VIDEO_DIR, audio_filename)

    if os.path.exists(video_path):
        print(f"🎥 Video đã sẵn sàng: {video_path}")
    else:
        print(f"❌ Lỗi! Video chưa được sẵn sàng tại {video_path}")
        return None
    if not transcribe.extract_audio(video_path, audio_path):
        return None

    if os.path.exists(audio_path):
        print(f"🎵 Âm thanh đã lưu tại: {audio_path}")
    else:
        print(f"❌ Lỗi! Âm thanh chưa được lưu vào {audio_path}")
        return None

    transcript = transcribe.transcribe_whisper(
        audio_path, use_chunks=config.USE_TITLE_MODE, max_chunks=config.MAX_CHUNKS_RUNTIME
    )
    transcript_filename = f"transcript_{video_number}.txt"
    transcript_path = os.path.join(config.OUTPUT_VIDEO_DIR, transcript_filename)
    with open(transcript_path, "w", encoding="utf-8") as f:
        if transcript is None:
            print("⚠ Không có transcript, bỏ qua xử lý văn bản.")
            transcript = ""
        formatted_transcript = transcribe.format_transcript_by_spaces(transcript, spaces_per_line=25)
        f.write(formatted_transcript)

    print(f"💾 Đã lưu transcript vào: {transcript_path}")

    try:
        if os.path.exists(audio_path):
            os.remove(audio_path)
            print(f"🗑️ Đã xóa file âm thanh: {audio_path}")
    except Exception as e:
        print(f"⚠️ Không thể xóa audio: {e}")

    if not is_local_file and os.path.exists(video_path):
        try:
            os.remove(video_path)
            print(f"🗑️ Đã xóa file video: {video_path}")
        except Exception as e:
            print(f"⚠️ Không thể xóa video: {e}")

    gc.collect()
    if config.device == "cuda":
        config.torch.cuda.empty_cache()

    if config.USE_TITLE_MODE:
        print("📁 Chế độ chia nhỏ: sẽ xử lý GPT sau khi nhận diện xong tất cả...")
        refined_text = None
        title_suggestion = None
    elif transcript:
        print(f"📝 Đang gửi nội dung video {video_number} đến ChatGPT...")
        refined_text = browser_client.send_to_chatgpt(driver, transcript)
        save_refined_text(refined_text, video_number)
        title_suggestion = None
    else:
        print(f"⚠️ Video {video_number} không có transcript, bỏ qua gửi đến ChatGPT.")
        refined_text = "Dây là video quá dài để lấy nội dung."
        title_suggestion = None

    para = doc.add_paragraph()
    run = para.add_run(f"{video_number}. {link}")
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.size = Pt(11)
    doc.add_paragraph(f"Tiêu đề: {title}")
    if title_suggestion:
        doc.add_paragraph(f"Tiêu đề gợi ý: {title_suggestion}")

    if transcript:
        if refined_text:
            doc.add_paragraph(f"Nội dung:\n{refined_text}")
        elif not config.USE_TITLE_MODE or (
            config.USE_TITLE_MODE and config.TITLE_MODE_SUBMODE == "content_only"
        ):
            doc.add_paragraph(f"Nội dung:\n{transcript}")
    else:
        doc.add_paragraph("Nội dung: (Video quá dài để lấy nội dung)")

    doc.add_paragraph("=" * 40)

    try:
        doc.save(config.OUTPUT_WORD_FILE)
        print("💾 Đã lưu kết quả thành công!")
    except Exception as e:
        print(f"❌ Lỗi khi lưu file Word: {e}")
    finally:
        print("✅ Hoàn thành xử lý video, tiếp tục với video tiếp theo!")


def process_videos():
    doc = Document()
    doc.save(config.OUTPUT_WORD_FILE)

    driver = browser_client.global_driver

    need_gpt = (not config.USE_TITLE_MODE) or (
        config.USE_TITLE_MODE and config.TITLE_MODE_SUBMODE != "content_only"
    )

    if need_gpt:
        with browser_client.DRIVER_LOCK:
            in_progress = browser_client.FIREFOX_INIT_IN_PROGRESS

        if (driver is None or not browser_client.is_driver_alive(driver)) and in_progress:
            print("⏳ Firefox đang khởi động (thread nền). Đang chờ...")
            browser_client.FIREFOX_INIT_EVENT.wait(timeout=90)
            driver = browser_client.global_driver

        if driver is None or not browser_client.is_driver_alive(driver):
            print("🌐 Chưa có Firefox/driver. Tự mở Firefox để dùng GPT...")
            driver = browser_client.init_firefox()
            browser_client.global_driver = driver

    with open(config.LINKS_FILE, "r", encoding="utf-8") as file:
        links = [line.strip() for line in file.readlines() if line.strip()]

    for idx, raw_line in enumerate(links, start=1):
        video_number, link = clean_link(raw_line)
        print(f"🔄 Đang xử lý video {idx}/{len(links)} (ID={video_number})...")

        if config.USE_TITLE_MODE:
            process_video(idx, raw_line, doc, None, None)
            transcript_file = os.path.join(config.OUTPUT_VIDEO_DIR, f"transcript_{video_number}.txt")
        else:
            process_video(idx, raw_line, doc, driver, driver)
            transcript_file = None

        if config.USE_TITLE_MODE and transcript_file and os.path.exists(transcript_file):
            with open(transcript_file, "r", encoding="utf-8") as f:
                current_transcript = f.read().strip()

                if current_transcript:
                    if config.TITLE_MODE_SUBMODE == "content_only":
                        print(f"\n📝 Chế độ Nội dung: đã lưu transcript cho video {idx}, bỏ qua gửi GPT.")
                        try:
                            os.remove(transcript_file)
                            print(f"🗑️ Đã xóa file transcript: {transcript_file}")
                        except Exception as e:
                            print(f"⚠️ Không thể xóa transcript: {e}")
                        continue

                    print(f"\n📝 Đang xử lý GPT cho video {idx}...")
                    print("🔄 Sử dụng Firefox đã được mở từ khi khởi động...")

                    print(f"\n🎯 Yêu cầu 1: Gợi ý 5 tiêu đề cho video {idx}...")
                    title_prompt = (
                        "Dựa trên nội dung video sau, hãy đề xuất 5 tiêu đề hấp dẫn, phù hợp cho kênh Haio6VietSub. "
                        "Mỗi tiêu đề trên một dòng, không ghi số thứ tự:\n\n" + current_transcript
                    )
                    titles_result = browser_client.send_to_chatgpt(driver, title_prompt, mode="raw")
                    if titles_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"🎯 5 Tiêu đề gợi ý cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(titles_result)

                    if config.TITLE_MODE_SUBMODE == "title_only":
                        doc.save(config.OUTPUT_WORD_FILE)
                        try:
                            os.remove(transcript_file)
                            print(f"🗑️ Đã xóa file transcript: {transcript_file}")
                        except Exception as e:
                            print(f"⚠️ Không thể xóa transcript: {e}")
                        print(f"✅ Video {idx}: Chế độ Tiêu đề - đã xong sau Yêu cầu 1.")
                        continue

                    print(f"\n📝 Yêu cầu 2: Viết mô tả video {idx} và hashtag...")
                    desc_prompt = (
                        "Viết mô tả video (1-2 đoạn ngắn gọn) và thêm 5-10 hashtag phù hợp cho kênh Haio6VietSub, "
                        "dựa trên nội dung:\n\n" + current_transcript
                    )
                    desc_result = browser_client.send_to_chatgpt(driver, desc_prompt, mode="raw")
                    if desc_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"📝 Mô tả + Hashtag cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(desc_result)

                    print(f"\n🏷️ Yêu cầu 3: Viết thẻ tag cho video {idx} (tối đa 500 kí tự)...")
                    tag_prompt = (
                        "Viết một danh sách các thẻ tag mô tả video này, tối đa 500 kí tự, phù hợp cho kênh Haio6VietSub, "
                        "ngăn cách bằng dấu phẩy. Dựa trên nội dung:\n\n" + current_transcript
                    )
                    tag_result = browser_client.send_to_chatgpt(driver, tag_prompt, mode="raw")
                    if tag_result:
                        para = doc.add_paragraph()
                        para_run = para.add_run(f"🏷️ Tags cho video {idx}:")
                        para_run.bold = True
                        doc.add_paragraph(tag_result)

                    doc.save(config.OUTPUT_WORD_FILE)

                    try:
                        os.remove(transcript_file)
                        print(f"🗑️ Đã xóa file transcript: {transcript_file}")
                    except Exception as e:
                        print(f"⚠️ Không thể xóa transcript: {e}")

                    print(f"✅ Đã lưu kết quả video {idx} vào {config.OUTPUT_WORD_FILE}")

                    gc.collect()
                    if config.device == "cuda":
                        config.torch.cuda.empty_cache()

    print("\n🧹 Giải phóng bộ nhớ...")

    if config.model is not None:
        del config.model
        config.model = None

    if config.model2 is not None:
        del config.model2
        config.model2 = None

    if config.device == "cuda":
        config.torch.cuda.empty_cache()
    gc.collect()

    print("✅ Bộ nhớ đã được giải phóng.")

    print("🎉 Hoàn thành tất cả video! Kết quả đã lưu vào final_result.docx")
    play_sound("success")
