import os
import re
import time
import sys

from selenium import webdriver
from selenium.webdriver.firefox.service import Service as FirefoxService
from selenium.webdriver.firefox.options import Options as FirefoxOptions
from selenium.webdriver.support.ui import WebDriverWait

from docx import Document

# ====== cấu hình đường dẫn ======
# Thêm thư mục chứa profiles_config.py vào sys.path
sys.path.append(r"D:\Python\GetLinktoText")
from profiles_config import FIREFOX_PROFILE_PATH  # trong đó bạn cấu hình profile Firefox

LINKS_FILE = r"D:\Python\GetLinktoText\linkgettoText.txt"          # file chứa link
OUTPUT_WORD_FILE = r"D:\Python\GetLinktoText\titles_only.docx"    # file Word để lưu tiêu đề
GECKODRIVER_PATH = r"D:\Python\GetLinktoText\geckodriver.exe"     # đường dẫn geckodriver


# ============ HÀM PHỤ ============

def clean_link(line, counter=[0]):
    """
    Nhận 1 dòng, có thể dạng:
      "1. https://..." hoặc "https://..."
    Trả về: (số thứ tự, link)
    """
    match = re.match(r'^\s*(\d+)[\.\s]+\s*(https?://\S+)', line)
    if match:
        counter[0] = int(match.group(1))
        return counter[0], match.group(2).strip()
    else:
        counter[0] += 1
        return counter[0], line.strip()


def init_firefox():
    """Khởi tạo Firefox với profile sẵn (đã đăng nhập nếu cần)."""
    options = FirefoxOptions()
    options.add_argument("-profile")
    options.add_argument(FIREFOX_PROFILE_PATH)

    service = FirefoxService(executable_path=GECKODRIVER_PATH)
    driver = webdriver.Firefox(service=service, options=options)
    # driver.minimize_window()  # nếu muốn thu nhỏ lại
    return driver


from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

def get_page_title(driver, url, timeout=5):
    """
    Lấy tiêu đề/caption Reel/TikTok bằng cách thử lần lượt danh sách XPath.
    """
    print(f"🔍 Đang mở: {url}")
    driver.get(url)

    xpaths = [
        '/html/body/div[1]/div/div[1]/div/div[3]/div/div/div[1]/div[1]/div/div/div[1]/div/div[2]/div/div/div/div/div/div/div[2]/div[1]/div/div/div[2]/div/div/div[2]/span/div',        
        '/html/body/div[1]/div[2]/div[2]/div/div[2]/div[1]/div[1]/div[2]/div[2]/div[1]/div/h1',
        '/html/body/div[1]/div[2]/div[2]/div/div[2]/div[1]/div[1]/div[2]/div[2]/div[1]/div/div[2]'
  
    ]

    for xpath in xpaths:
        try:
            element = WebDriverWait(driver, timeout).until(
                EC.presence_of_element_located((By.XPATH, xpath))
            )
            text = element.text.strip()
            if text:
                print(f"✅ Tiêu đề (theo XPath): {text}")
                return text
        except:
            print(f"⚠️ Không tìm thấy tiêu đề bằng XPath: {xpath}")

    print("❌ Không lấy được tiêu đề — có thể Reel không load caption hoặc không có caption.")
    return "Không thể lấy tiêu đề"

# ============ HÀM CHÍNH ============

def process_links():
    # Xóa trắng file Word trước khi ghi mới
    from docx import Document
    doc = Document()  
    doc.save(OUTPUT_WORD_FILE)
    print("🗑️ Đã xóa trắng file titles_only.docx")

    if not os.path.exists(LINKS_FILE):
        print(f"❌ Không tìm thấy file {LINKS_FILE}")
        return

    # Đọc danh sách link
    with open(LINKS_FILE, "r", encoding="utf-8") as f:
        raw_lines = [line.strip() for line in f.readlines() if line.strip()]

    if not raw_lines:
        print("⚠️ File link trống, không có gì để xử lý.")
        return

    print(f"📄 Đã đọc {len(raw_lines)} dòng link từ {LINKS_FILE}")

    # Chuẩn bị file Word: nếu có thì mở, không có thì tạo mới
    if os.path.exists(OUTPUT_WORD_FILE):
        doc = Document(OUTPUT_WORD_FILE)
        # Nếu muốn xóa trắng nội dung cũ mỗi lần chạy:
        # doc = Document()
    else:
        doc = Document()

    driver = init_firefox()

    try:
        for idx, line in enumerate(raw_lines, start=1):
            stt, url = clean_link(line)

            # Bỏ qua nếu dòng không phải link http
            if not url.lower().startswith("http"):
                print(f"⚠️ Bỏ qua dòng {stt}: không phải link hợp lệ ({url})")
                continue

            print(f"\n🎥 Đang xử lý link {stt}: {url}")
            title = get_page_title(driver, url)

            # Mỗi link → 1 dòng trong Word: "1. Facebook"
            para = doc.add_paragraph(f"{stt}. {title}")

            # Lưu luôn sau mỗi link (để có gì lỗi giữa chừng vẫn có dữ liệu)
            doc.save(OUTPUT_WORD_FILE)
            print(f"💾 Đã ghi vào {OUTPUT_WORD_FILE}: {stt}. {title}")

        print(f"\n🎉 Hoàn thành! Đã lưu tiêu đề vào file Word: {OUTPUT_WORD_FILE}")

    finally:
        driver.quit()
        print("🛑 Đã đóng Firefox.")


if __name__ == "__main__":
    process_links()
