import os
import time
import yt_dlp
import whisper
import torch
import subprocess
import shutil
import wave
import re
from docx import Document
from docx.shared import RGBColor, Pt
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
import re

# ================= Cấu hình =================
LINKS_FILE = "linkgettoText.txt"  # File chứa danh sách link video
OUTPUT_WORD_FILE = "final_result.docx"  # Xuất kết quả ra file Word
WHISPER_MODEL_PATH = "whisper_cache"  # Thư mục lưu mô hình Whisper
OUTPUT_VIDEO_DIR = "output_videos"  # Thư mục lưu video và âm thanh
MODEL_NAME = "large-v3"  # Mô hình Whisper
FFMPEG_PATH = shutil.which("ffmpeg")  # Kiểm tra FFmpeg

# Danh sách các tên bác sĩ cần thay thế
NAME_VARIANTS = [
    "bác sĩ A", "bác sĩ B", "bác sĩ C",  # Thêm các tên bác sĩ không phải "bác sĩ Hội"
    "BS. A", "BS. B", "BS. C",  # Các cách viết tắt có thể xuất hiện
    "BS A", "BS B", "BS C"
]

# Kiểm tra FFmpeg
if FFMPEG_PATH is None:
    print("⚠️ FFmpeg không có trong PATH. Hãy cài đặt FFmpeg để tiếp tục.")
    exit(1)

# Kiểm tra GPU
device = "cuda" if torch.cuda.is_available() else "cpu"
torch_dtype = torch.float16 if device == "cuda" else torch.float32
print(f"✅ Whisper sẽ chạy trên {'GPU' if device == 'cuda' else 'CPU'}")

# Tạo thư mục nếu chưa có
os.makedirs(WHISPER_MODEL_PATH, exist_ok=True)
os.makedirs(OUTPUT_VIDEO_DIR, exist_ok=True)

# Tải mô hình Whisper
model = whisper.load_model(MODEL_NAME, download_root=WHISPER_MODEL_PATH).to(device)

# Xóa tất cả file trong thư mục nhưng giữ lại thư mục
def clear_output_directory(directory):
    if os.path.exists(directory):
        for file in os.listdir(directory):
            file_path = os.path.join(directory, file)
            try:
                if os.path.isfile(file_path) or os.path.islink(file_path):
                    os.unlink(file_path)  # Xóa file hoặc symbolic link
                elif os.path.isdir(file_path):
                    shutil.rmtree(file_path)  # Xóa thư mục con nếu có
            except Exception as e:
                print(f"⚠️ Lỗi khi xóa {file_path}: {e}")

# Gọi hàm xóa file trong thư mục output_video trước khi bắt đầu xử lý
clear_output_directory(OUTPUT_VIDEO_DIR)

def init_browser():
    """Khởi tạo trình duyệt Selenium."""
    chrome_options = Options()
    chrome_options.add_argument("--headless")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--window-size=1920x1080")
    service = Service(ChromeDriverManager().install())
    return webdriver.Chrome(service=service, options=chrome_options)

def get_video_title(url):
    """Lấy tiêu đề video bằng Selenium."""
    browser = init_browser()
    try:
        browser.get(url)
        wait = WebDriverWait(browser, 10)
        xpath = '//h1'  # XPath có thể thay đổi
        title_element = wait.until(EC.presence_of_element_located((By.XPATH, xpath)))
        title = title_element.text.strip()
        print(f"✅ Tiêu đề video: {title}")
        return title
    except Exception as e:
        print(f"❌ Lỗi khi lấy tiêu đề: {e}")
        return "Không thể lấy tiêu đề"
    finally:
        browser.quit()

def clean_link(line):
    """Loại bỏ số thứ tự trước link."""
    match = re.search(r'(\d+)\.\s*(https?://\S+)', line)
    return (int(match.group(1)), match.group(2)) if match else (None, line.strip())

def replace_doctor_names(text):
    """Thay thế tất cả tên bác sĩ khác thành 'bác sĩ Hội'."""
    for name in NAME_VARIANTS:
        text = re.sub(rf"\b{name}\b", "bác sĩ Hội", text, flags=re.IGNORECASE)
    return text

def download_video(video_url, video_path):
    """Tải video từ TikTok hoặc YouTube"""
    ydl_opts = {
        "outtmpl": video_path,
        "format": "bestvideo+bestaudio/best",
        "merge_output_format": "mp4",
        "quiet": True,
        "noplaylist": True,
    }
    try:
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            ydl.download([video_url])
        print(f"✅ Video đã tải thành công: {video_path}")
        return video_path
    except Exception as e:
        print(f"❌ Lỗi khi tải video: {e}")
        return None

def extract_audio(video_path, audio_path):
    """Trích xuất âm thanh từ video"""
    command = [FFMPEG_PATH, "-i", video_path, "-ac", "1", "-ar", "16000", "-y", audio_path, "-vn"]
    try:
        subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)
        print(f"✅ Âm thanh đã trích xuất thành công: {audio_path}")
        return audio_path
    except subprocess.CalledProcessError as e:
        print(f"❌ Lỗi khi trích xuất âm thanh: {e}")
        return None

def check_audio_file(audio_path):
    """Kiểm tra file âm thanh có hợp lệ không"""
    try:
        with wave.open(audio_path, "rb") as f:
            frames = f.getnframes()
            rate = f.getframerate()
            duration = frames / float(rate)
            print(f"🎵 Thời lượng file âm thanh: {duration:.2f} giây")
            return duration > 0
    except Exception as e:
        print(f"❌ Lỗi khi kiểm tra file âm thanh: {e}")
        return False

def transcribe_audio(audio_path):
    """Nhận diện giọng nói bằng Whisper"""
    if not check_audio_file(audio_path):
        print("⚠️ File âm thanh không hợp lệ, dừng xử lý.")
        return None

    print("📝 Đang nhận diện giọng nói...")
    result = model.transcribe(audio_path, language="vi")
    cleaned_text = replace_doctor_names(result["text"])
    return cleaned_text

def process_video(index, link, doc):
    """Xử lý một video và ghi vào file Word"""
    video_number, link = clean_link(link)
    if video_number is None:
        video_number = index

    print(f"\n🎥 Đang xử lý video {video_number}: {link}")
    title = get_video_title(link)

    video_filename = f"video_{video_number}.mp4"
    audio_filename = f"audio_{video_number}.wav"
    video_path = os.path.join(OUTPUT_VIDEO_DIR, video_filename)
    audio_path = os.path.join(OUTPUT_VIDEO_DIR, audio_filename)

    downloaded_video = download_video(link, video_path)
    if not downloaded_video:
        return None
    
    extracted_audio = extract_audio(downloaded_video, audio_path)
    if not extracted_audio:
        return None
    
    transcript = transcribe_audio(extracted_audio)
    
    link_paragraph = doc.add_paragraph()
    link_run = link_paragraph.add_run(f"{video_number}. {link}")
    link_run.font.color.rgb = RGBColor(255, 0, 0)
    link_run.font.size = Pt(12)

    title_paragraph = doc.add_paragraph()
    title_run = title_paragraph.add_run("Tiêu đề: ")
    title_run.bold = True
    title_paragraph.add_run(title)

    content_paragraph = doc.add_paragraph()
    content_run = content_paragraph.add_run("Nội dung:\n")
    content_run.bold = True
    content_paragraph.add_run(transcript)

    doc.add_paragraph("=" * 40)

def process_videos():
    """Xử lý tất cả video và lưu vào file Word"""
    doc = Document()
    with open(LINKS_FILE, "r", encoding="utf-8") as file:
        links = [line.strip() for line in file.readlines() if line.strip()]
    
    for idx, link in enumerate(links, start=1):
        process_video(idx, link, doc)
    
    doc.save(OUTPUT_WORD_FILE)
    print(f"🎉 Hoàn thành tất cả video! Kết quả đã lưu vào {OUTPUT_WORD_FILE}")

if __name__ == "__main__":
    process_videos()
