import os
import re
import time
from selenium import webdriver
from selenium.webdriver.firefox.service import Service as FirefoxService
from selenium.webdriver.firefox.options import Options as FirefoxOptions
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.firefox import GeckoDriverManager
from docx import Document  
from docx.shared import RGBColor

# ================= Cấu hình =================
OUTPUT_VIDEO_DIR = "output_videos"  # Thư mục chứa file transcript_xx.txt textRender2
OUTPUT_WORD_FILE = "final_result.docx"  # File Word xuất kết quả
FIREFOX_PROFILE_PATH = r"C:\Users\PC\AppData\Roaming\Mozilla\Firefox\Profiles\kt92li14.GPT_DEEPSEEK"

def init_firefox():
    """Khởi tạo trình duyệt Firefox với Selenium"""
    options = FirefoxOptions()
    options.add_argument(f"-profile")
    options.add_argument(f"{FIREFOX_PROFILE_PATH}")
    service = FirefoxService(GeckoDriverManager().install())
    driver = webdriver.Firefox(service=service, options=options)
    driver.minimize_window()
    driver.get("https://chat.openai.com/")
    time.sleep(7)
    return driver

def send_to_chatgpt(driver, text):
    """Gửi nội dung đến ChatGPT và lấy phản hồi"""
    try:
        prompt = f"trả lời nội dung không cần viết lại yêu cầu, sửa kiểm chính tả cho văn bản tiếng việt kèm 1 số từ tiếng anh nếu có, kiểm tra câu văn có nghĩa, viết thành 1 đoạn duy nhất, thêm dấu câu, thay Bác sĩ x thành bác sĩ Hội:\n\n{text}"

        textarea = WebDriverWait(driver, 20).until(
            EC.presence_of_element_located((By.XPATH, '//*[@id="prompt-textarea"]'))
        )
        
        textarea.clear()
        textarea.send_keys(prompt)
        textarea.send_keys("\ue007")  # Nhấn Enter

        max_wait_time = 180  # Chờ tối đa 3 phút
        check_interval = 5  
        stable_time = 30  
        stop_button_xpath = "//button[@data-testid='stop-button']"  

        print("⌛ Đang chờ phản hồi từ ChatGPT...")

        start_time = time.time()
        chatgpt_started_typing = False

        while time.time() - start_time < max_wait_time:
            time.sleep(check_interval)
            stop_button = driver.find_elements(By.XPATH, stop_button_xpath)
            if stop_button:
                print("✅ ChatGPT bắt đầu phản hồi!")
                chatgpt_started_typing = True
                break

        if not chatgpt_started_typing:
            print("❌ ChatGPT không phản hồi, bỏ qua.")
            return None

        while True:
            time.sleep(check_interval)
            stop_button = driver.find_elements(By.XPATH, stop_button_xpath)
            if not stop_button:
                print("✅ ChatGPT đã dừng gõ.")
                break

        messages = driver.find_elements(By.CSS_SELECTOR, "div[data-message-author-role='assistant']")
        final_content = "\n".join([msg.text.strip() for msg in messages])

        if final_content:
            print("✅ Nhận được phản hồi từ GPT.")
            return final_content
        else:
            print("❌ Nội dung trống.")
            return None

    except Exception as e:
        print(f"❌ Lỗi khi gửi đến ChatGPT: {e}")
        return None

def process_transcripts():
    """Xử lý các file transcript, gửi lên ChatGPT và lưu vào Word"""
    
    if os.path.exists(OUTPUT_WORD_FILE):
        doc = Document(OUTPUT_WORD_FILE)
    else:
        doc = Document()
        doc.save(OUTPUT_WORD_FILE)

    driver = init_firefox()
    
    transcript_files = sorted([f for f in os.listdir(OUTPUT_VIDEO_DIR) if f.startswith("transcript_") and f.endswith(".txt")])
    
    for transcript_file in transcript_files:
        video_number = re.search(r"transcript_(\d+)\.txt", transcript_file)
        if not video_number:
            continue
        video_number = video_number.group(1)

        transcript_path = os.path.join(OUTPUT_VIDEO_DIR, transcript_file)
        with open(transcript_path, "r", encoding="utf-8") as file:
            transcript = file.read().strip()

        if not transcript:
            print(f"⚠ Bỏ qua video {video_number} vì không có nội dung.")
            continue

        print(f"📩 Đang gửi nội dung video {video_number} lên ChatGPT...")
        refined_text = send_to_chatgpt(driver, transcript)

        if refined_text:
            doc.add_paragraph(f"{video_number}. {transcript_file}").runs[0].font.color.rgb = RGBColor(255, 0, 0)
            doc.add_paragraph(f"Nội dung:\n{refined_text}")
            doc.add_paragraph("=" * 40)

            try:
                doc.save(OUTPUT_WORD_FILE)
                print(f"💾 Đã lưu kết quả của video {video_number} vào Word.")
            except Exception as e:
                print(f"❌ Lỗi khi lưu file Word: {e}")

    driver.quit()
    print("🎉 Hoàn thành việc xử lý tất cả các transcript!")

if __name__ == "__main__":
    process_transcripts()
